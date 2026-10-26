#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""完整剧本 Word 生成器（编剧工坊 host 插件调用）。

输入：demo/剧本_第01集.json
输出：可直接打印/围读的剧本 docx —— 封面、人物表、正文、页眉页脚。

格式遵循悉德·菲尔德《电影剧本写作基础》的剧本规范，并适配中文排版：
    场次标题（加粗）  动作描述（宋体）  人物名（居中）
    括号提示（居中、斜体）  对白（左右缩进）  转场（右对齐）
"""

from __future__ import annotations

import argparse
import json
import os
import sys

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

INK = RGBColor(0x1B, 0x2A, 0x38)
GREY = RGBColor(0x6B, 0x72, 0x80)
TEAL = RGBColor(0x1F, 0x6F, 0x78)
AMBER = RGBColor(0x9A, 0x6B, 0x1F)

FONT_CN = "宋体"
FONT_HEAD = "微软雅黑"
# macOS / Linux 没有微软雅黑与宋体：换成系统自带的宋体/黑体，
# 并允许用环境变量强制指定（SCREENWRITER_FONT_SANS / SCREENWRITER_FONT_SERIF）。
if sys.platform == "darwin":
    FONT_CN = "Songti SC"
    FONT_HEAD = "PingFang SC"
elif sys.platform.startswith("linux"):
    FONT_CN = "Noto Serif CJK SC"
    FONT_HEAD = "Noto Sans CJK SC"
FONT_CN = os.environ.get("SCREENWRITER_FONT_SERIF", FONT_CN)
FONT_HEAD = os.environ.get("SCREENWRITER_FONT_SANS", FONT_HEAD)


def set_run_font(run, name=FONT_CN, size=12, bold=False, italic=False, color=INK):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)
    rfonts.set(qn("w:eastAsia"), name)


def para(doc, text="", style=None, align=None, before=0, after=0, line=None,
         first_indent=None, left_indent=None, right_indent=None,
         font=FONT_CN, size=12, bold=False, italic=False, color=INK):
    p = doc.add_paragraph(style=style)
    if align is not None:
        p.alignment = align
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    if line is not None:
        pf.line_spacing = line
    if first_indent is not None:
        pf.first_line_indent = first_indent
    if left_indent is not None:
        pf.left_indent = left_indent
    if right_indent is not None:
        pf.right_indent = right_indent
    if text:
        run = p.add_run(text)
        set_run_font(run, font, size, bold, italic, color)
    return p


def add_page_number(paragraph):
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_end)
    set_run_font(run, FONT_HEAD, 9, color=GREY)


def shade(cell, hex_color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tc_pr.append(shd)


def set_cell_margins(table, top=60, bottom=60, left=100, right=100):
    tbl_pr = table._tbl.tblPr
    mar = OxmlElement("w:tblCellMar")
    for side, val in (("top", top), ("bottom", bottom), ("left", left), ("right", right)):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:w"), str(val))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tbl_pr.append(mar)


def build(data, out_path):
    doc = Document()

    # 页面
    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(2.4)
    sec.bottom_margin = Cm(2.2)
    sec.left_margin = Cm(2.8)
    sec.right_margin = Cm(2.4)

    # 正文默认样式
    normal = doc.styles["Normal"]
    normal.font.name = FONT_CN
    normal.font.size = Pt(12)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)

    # 页眉
    header = sec.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_run_font(hp.add_run(f"{data['title']} · {data['episode']} · {data['draft']}"), FONT_HEAD, 8.5, color=GREY)

    # 页脚页码
    footer = sec.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_page_number(fp)

    # ---------------------------------------------------------------- 封面
    para(doc, "编剧工坊 · 剧本交付", align=WD_ALIGN_PARAGRAPH.LEFT, before=30, after=6,
         font=FONT_HEAD, size=9, bold=True, color=TEAL)

    para(doc, data["title"], align=WD_ALIGN_PARAGRAPH.LEFT, before=10, after=2,
         font=FONT_HEAD, size=40, bold=True, color=INK)
    para(doc, f"{data['episode']}  {data['episode_title']}", align=WD_ALIGN_PARAGRAPH.LEFT, before=0, after=10,
         font=FONT_HEAD, size=18, color=TEAL)

    para(doc, data["format"], align=WD_ALIGN_PARAGRAPH.LEFT, after=2, font=FONT_CN, size=11, color=GREY)
    para(doc, f"{data['draft']}    {data['date']}", align=WD_ALIGN_PARAGRAPH.LEFT, after=24, font=FONT_CN, size=11, color=GREY)

    # 一句话梗概
    para(doc, "一句话梗概 / LOGLINE", before=6, after=4, font=FONT_HEAD, size=9, bold=True, color=AMBER)
    para(doc, data["logline"], after=24, size=12, color=INK, line=1.6)

    # 人物表（两列表格）
    para(doc, "人物表 / CHARACTERS", before=6, after=6, font=FONT_HEAD, size=9, bold=True, color=AMBER)
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    set_cell_margins(table)
    hdr = table.rows[0].cells
    for i, text in enumerate(["人物", "人物小传"]):
        hdr[i].text = ""
        p = hdr[i].paragraphs[0]
        run = p.add_run(text)
        set_run_font(run, FONT_HEAD, 9, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
        shade(hdr[i], "1B2A38")
    for ch in data["characters"]:
        cells = table.add_row().cells
        for i, text in enumerate([ch["name"], ch["desc"]]):
            cells[i].text = ""
            p = cells[i].paragraphs[0]
            run = p.add_run(text)
            set_run_font(run, FONT_CN, 10, bold=(i == 0), color=INK)
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
        shade(cells[0], "F3ECDD")
    for row in table.rows:
        row.cells[0].width = Cm(3.0)
        row.cells[1].width = Cm(12.5)

    para(doc, "本稿为 demo：结构与格式完整，个别场面留待下一稿精修。", before=20, after=0,
         font=FONT_CN, size=9, italic=True, color=GREY)

    # ---------------------------------------------------------------- 正文
    p = para(doc, "正 文", align=WD_ALIGN_PARAGRAPH.CENTER, before=0, after=14,
             font=FONT_HEAD, size=14, bold=True, color=INK)
    p.paragraph_format.page_break_before = True

    for scene in data["scenes"]:
        # 场次标题
        p = para(doc, scene["heading"], before=18, after=8, font=FONT_HEAD, size=12, bold=True, color=INK)
        p.paragraph_format.keep_with_next = True

        for el in scene["elements"]:
            t = el.get("type")
            text = el.get("text", "")
            if t == "action":
                para(doc, text, before=2, after=2, line=1.55, first_indent=Cm(0.0))
            elif t == "character":
                p = para(doc, text, align=WD_ALIGN_PARAGRAPH.CENTER, before=8, after=0,
                         font=FONT_CN, size=11.5, bold=True, color=INK)
                p.paragraph_format.keep_with_next = True
            elif t == "parenthetical":
                p = para(doc, text, align=WD_ALIGN_PARAGRAPH.CENTER, before=0, after=0,
                         font=FONT_CN, size=11, italic=True, color=GREY)
                p.paragraph_format.keep_with_next = True
            elif t == "dialogue":
                para(doc, text, align=WD_ALIGN_PARAGRAPH.LEFT, before=0, after=6, line=1.45,
                     left_indent=Cm(3.2), right_indent=Cm(2.6), size=12)
            elif t == "transition":
                para(doc, text, align=WD_ALIGN_PARAGRAPH.RIGHT, before=8, after=8,
                     font=FONT_HEAD, size=10.5, bold=True, color=TEAL)

    # 尾页
    p = para(doc, "— 第 1 集 完 —", align=WD_ALIGN_PARAGRAPH.CENTER, before=80, after=20,
             font=FONT_HEAD, size=16, bold=True, color=INK)
    p.paragraph_format.page_break_before = True
    para(doc, "下一集：《杂音》。林晚开始修复磁带，发现第二个人声属于一个十八岁的少年。",
         align=WD_ALIGN_PARAGRAPH.CENTER, before=6, after=6, font=FONT_CN, size=10.5, color=GREY)
    para(doc, "《回声镇》· 编剧工坊 demo · 2026",
         align=WD_ALIGN_PARAGRAPH.CENTER, before=30, font=FONT_CN, size=9, color=GREY)

    os.makedirs(os.path.dirname(os.path.abspath(out_path)), exist_ok=True)
    doc.core_properties.title = f"{data['title']} {data['episode']} {data['episode_title']}"
    doc.core_properties.author = "编剧工坊"
    doc.save(out_path)
    return out_path


def main(argv=None):
    ap = argparse.ArgumentParser()
    ap.add_argument("--script", required=True, help="剧本 JSON")
    ap.add_argument("--out", required=True, help="输出 docx 路径")
    args = ap.parse_args(argv)
    with open(args.script, "r", encoding="utf-8") as f:
        data = json.load(f)
    out = build(data, args.out)
    print(json.dumps({"ok": True, "out": os.path.abspath(out)}, ensure_ascii=True))
    return 0


if __name__ == "__main__":
    sys.exit(main())
