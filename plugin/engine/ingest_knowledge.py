#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""把上传的剧本/编剧书抽成可检索的 Markdown 摘录。

取字顺序：
  .docx  → python-docx
  .pdf   → PyMuPDF 读文字层；没有文字层的扫描页交给本机 RapidOCR
  .txt / .md → 原样保留
OCR 依赖复用 course-kit/.pylibs（RapidOCR + pymupdf），不需要联网。
"""

from __future__ import annotations

import argparse
import json
import os
import sys
from pathlib import Path

MAX_CHARS = 300_000


def ensure_libs() -> None:
    """确保 course-kit/.pylibs 在 sys.path（docx 也装在那里）。"""
    sys.path.insert(0, str(Path(__file__).resolve().parent))
    try:
        import ocr_pdf  # noqa: PLC0415

        ocr_pdf._ensure_libs()
    except Exception:
        pass


def extract_docx(path: Path) -> str:
    from docx import Document

    doc = Document(str(path))
    parts = [p.text.strip() for p in doc.paragraphs if p.text and p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text and c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n\n".join(parts)


def extract_pdf(path: Path) -> tuple[str, str, dict]:
    """返回 (正文, 方法, 统计)。文字层优先；扫描页才 OCR。"""
    sys.path.insert(0, str(Path(__file__).resolve().parent))
    import ocr_pdf  # noqa: E402  同目录脚本

    limit = int(os.environ.get("SCREENWRITER_OCR_PAGES") or ocr_pdf.DEFAULT_PAGES)
    result = ocr_pdf.extract_text(str(path), limit)
    if result.get("ocrPages"):
        method = "pymupdf+rapidocr"
    else:
        method = "pymupdf"
    detail = {k: v for k, v in result.items() if k != "text"}
    return result.get("text", ""), method, detail


def main(argv=None) -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--src", required=True, help="上传的文件")
    ap.add_argument("--outDir", required=True, help="摘录输出目录")
    args = ap.parse_args(argv)

    src = Path(args.src)
    out_dir = Path(args.outDir)
    ext = src.suffix.lower()
    text = ""
    method = ""
    detail: dict = {}
    ensure_libs()

    try:
        if ext == ".docx":
            text, method = extract_docx(src), "python-docx"
        elif ext == ".pdf":
            text, method, detail = extract_pdf(src)
        elif ext in (".txt", ".md"):
            text, method = src.read_text(encoding="utf-8", errors="ignore"), "text"
        else:
            method = "unsupported"
    except Exception as e:  # noqa: BLE001
        print(json.dumps({"ok": False, "message": str(e)}, ensure_ascii=True))
        return 1

    ocr_note = ""
    if detail.get("ocrPages"):
        ocr_note = (
            f"其中 {detail['ocrPages']} 页是扫描页，已用 RapidOCR 本地识别"
            + (f"；另有 {detail['ocrSkipped']} 页超出单次 OCR 上限（{detail['ocrLimit']} 页）未识别"
               if detail.get("ocrSkipped") else "")
        )

    if not text.strip():
        print(json.dumps({
            "ok": False,
            "method": method,
            "chars": 0,
            "ocr": detail,
            "message": "没有抽到文字（可能是扫描件/图片 PDF）；原文件已保存。"
                       + ("OCR 也没认出字，试试提高清晰度或换文件。" if detail.get("ocrPages") else ""),
        }, ensure_ascii=True))
        return 0

    text = text[:MAX_CHARS]
    if ext in (".txt", ".md"):
        # 原文本身可读，不再生成重复摘录
        print(json.dumps({"ok": True, "method": method, "chars": len(text), "out": None}, ensure_ascii=True))
        return 0

    out_dir.mkdir(parents=True, exist_ok=True)
    out = out_dir / f"{src.stem}.摘录.md"
    header = (
        f"# 知识库摘录 · {src.name}\n\n"
        f"> 由编剧工坊自动抽取（{method}），{len(text)} 字；"
        + (f"{ocr_note}；" if ocr_note else "")
        + f"原文见同目录 `{src.name}`。\n\n"
    )
    out.write_text(header + text + "\n", encoding="utf-8")
    print(json.dumps({
        "ok": True,
        "method": method,
        "chars": len(text),
        "out": str(out),
        "ocr": detail or None,
    }, ensure_ascii=True))
    return 0


if __name__ == "__main__":
    sys.exit(main())
